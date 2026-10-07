"""Reading uploaded spreadsheets and guessing which columns hold what."""

import csv
import io
import re

from openpyxl import Workbook, load_workbook

NONE = "(not in my sheet)"


def read_upload(data, name):
    """Return (workbook to write into, workbook of values to read from)."""
    lower = name.lower()
    if lower.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        wb = Workbook()
        for row in csv.reader(io.StringIO(text), dialect):
            wb.active.append(row)
        return wb, wb
    if lower.endswith(".xls"):
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        wb = Workbook()
        wb.remove(wb.active)
        for sh in book.sheets():
            ws = wb.create_sheet(sh.name[:31])
            for r in range(sh.nrows):
                ws.append(sh.row_values(r))
        return wb, wb
    keep_vba = lower.endswith(".xlsm")
    return (load_workbook(io.BytesIO(data), keep_vba=keep_vba),
            load_workbook(io.BytesIO(data), data_only=True))


def guess_header_row(ws):
    """Headers are usually row 1, but some sheets start with a title. Take the first row in the
    top ten that mentions a company/name/number column, else the first row with text in it."""
    first = None
    for r in range(1, min(ws.max_row, 10) + 1):
        cells = [str(c.value).strip().lower() for c in ws[r] if c.value not in (None, "")]
        if cells and first is None:
            first = r
        if any(re.search(r"company|name|number|crn|reg", c) for c in cells):
            return r
    return first or 1


def guess(headers, keywords, skip=None):
    for kw in keywords:
        for h in headers:
            if h != skip and kw in h.lower():
                return h
    return NONE
