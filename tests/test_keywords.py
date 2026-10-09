"""Searching each company's accounts for a list of keywords."""

import io
from pathlib import Path

from openpyxl import Workbook

import turnover as t
from conftest import FakeWeb, company, filing_page, fixture, needs_ocr
from inputs import read_upload

DEFAULTS = t.parse_keywords((Path(t.__file__).parent / "default_keywords.txt").read_text())


def found(words, text, ocr=False):
    k = t.Keywords(words)
    k.scan(text, ocr)
    return k.result()


# ---------------------------------------------------------------- matching rules


def test_case_and_punctuation_dont_matter():
    assert found(["Foreign Exchange"], "the group's FOREIGN-EXCHANGE exposure") == ["Foreign Exchange"]


def test_phrases_survive_line_breaks_and_hyphenation():
    assert found(["foreign exchange"], "risks from foreign\nexchange movements") == ["foreign exchange"]
    assert found(["foreign exchange"], "risks from foreign ex-\nchange movements") == ["foreign exchange"]


def test_phrases_survive_ocr_running_words_together():
    text = "Thegroupmanagesitsforeignexchangeriskthroughforwardcontracts"
    assert found(["foreign exchange risk", "forward contracts", "treasury"], text, ocr=True) == \
        ["foreign exchange risk", "forward contracts"]
    assert found(["foreign exchange risk"], text) == []        # clean text needs real words


def test_short_words_must_stand_alone():
    assert found(["FX"], "fixed assets, FIXTURES and fittings") == []
    assert found(["FX"], "FX forwards") == ["FX"]
    assert found(["import"], "an important customer") == []
    assert found(["import"], "goods we import from Spain") == ["import"]
    assert found(["options"], "share options granted") == ["options"]


def test_results_keep_the_given_order_without_duplicates():
    assert found(["exports", "FX", "fx", "  ", "hedging"], "hedging of exports and FX") == ["exports", "FX", "hedging"]


def test_mentions_are_counted_and_ranked():
    k = t.Keywords(["hedging", "foreign exchange", "FX"])
    k.scan("FX risk. Foreign exchange forwards. FX swaps and foreign\nexchange options. FX")
    k.scan("Foreignexchange again on the next page", ocr=True)        # counts add up across pages
    assert k.ranked() == {"foreign exchange": 3, "FX": 3}               # tie keeps the given order
    assert k.result() == ["foreign exchange", "FX"]


def test_keyword_cell_shows_counts_most_mentioned_first():
    res = t.Result(keywords_found=["FX", "hedging"], keyword_counts={"hedging": 5, "FX": 2})
    assert t.keyword_cell(res) == "hedging (5), FX (2)" and t.keyword_count(res) == 2
    assert t.keyword_cell(t.Result(keywords_found=[])) == "none found" and t.keyword_count(t.Result(keywords_found=[])) == 0
    assert t.keyword_cell(t.Result()) is None and t.keyword_count(t.Result()) is None


def test_nothing_found_is_an_empty_list():
    assert found(["derivatives"], "a small bakery in Leeds") == []


def test_parse_keywords_accepts_lines_commas_and_semicolons():
    assert t.parse_keywords("FX\nhedging, derivatives;\n\n  treasury  ") == ["FX", "hedging", "derivatives", "treasury"]
    assert t.parse_keywords("") == [] and t.parse_keywords(None) == []


def test_default_keyword_list_loads():
    assert len(DEFAULTS) == 162 and "foreign exchange" in DEFAULTS and "FX" in DEFAULTS
    assert len(t.Keywords(DEFAULTS).words) == 162                 # no duplicates in the list


# ---------------------------------------------------------------- documents


def test_real_ixbrl_filing_is_searched():
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(fixture("ixbrl_tagged.html", "rb"), "html.parser")
    t.read_ixbrl(soup)
    k = t.Keywords(DEFAULTS)
    k.scan(t.ixbrl_text(soup))
    assert k.result() == ["derivative", "presentation currency", "derivative financial instruments"]


def test_hidden_ixbrl_header_is_not_searched():
    from bs4 import BeautifulSoup
    soup = BeautifulSoup('<html><body><ix:header><ix:hidden>foreign exchange</ix:hidden></ix:header>'
                         '<p>Turnover</p></body></html>', "html.parser")
    assert found(["foreign exchange"], t.ixbrl_text(soup)) == []


def test_text_pdf_is_searched_without_ocr(monkeypatch):
    monkeypatch.setattr(t, "ocr_engine", lambda: (_ for _ in ()).throw(AssertionError("no OCR for text PDFs")))
    k = t.Keywords(["gross profit", "hedging"])
    got = t.read_pdf(fixture("text_small_company.pdf", "rb"), keywords=k)
    assert got.turnover == 5_432_100 and k.result() == ["gross profit"] and got.keyword_pages == 1


@needs_ocr
def test_scanned_pdf_every_page_is_searched_and_turnover_still_found():
    k = t.Keywords(["joint ventures", "annual report", "cost of sales", "hedging"])
    got = t.read_pdf(fixture("scanned_annual_report.pdf", "rb"), keywords=k)
    assert got.turnover == 1_234_500_000 and got.page == 3
    assert k.result() == ["joint ventures", "annual report", "cost of sales"] and got.keyword_pages == 3


@needs_ocr
def test_every_page_is_read_so_counts_are_complete():
    k = t.Keywords(["joint ventures", "revenue"])
    got = t.read_pdf(fixture("scanned_annual_report.pdf", "rb"), keywords=k)
    assert got.keyword_pages == 3 and got.turnover == 1_234_500_000
    assert k.ranked() == {"revenue": 2, "joint ventures": 1}           # revenue on pages 2 and 3


def test_scanned_pages_can_be_left_out_of_the_search():
    k = t.Keywords(["joint ventures"])
    got = t.read_pdf(fixture("scanned_annual_report.pdf", "rb"), keywords=k, ocr_keywords=False, read_scanned=False)
    assert k.result() == [] and got.keyword_pages == 0


# ---------------------------------------------------------------- per company


def run(routes, **kw):
    return t.safe_check({"number": "01234567"}, FakeWeb(routes), **kw)


TAGGED = filing_page(("01 Mar 2026", "AA", "Full accounts made up to 31 December 2025", True, True))
MICRO = filing_page(("29 Jul 2026", "AA", "Micro company accounts made up to 31 October 2025", True, True))
SCANNED = filing_page(("25 Jul 2026", "AA", "Group of companies' accounts made up to 28 February 2026", True, False))


def test_company_with_tagged_accounts_gets_turnover_and_keywords():
    res = run(company("01234567", filings=TAGGED, xhtml=fixture("ixbrl_tagged.html", "rb")), keywords=DEFAULTS)
    assert res.status == "FOUND" and res.turnover == 44_043_650
    assert res.keywords_found == ["derivative", "presentation currency", "derivative financial instruments"]


def test_micro_company_searched_but_nothing_found():
    res = run(company("01234567", filings=MICRO, xhtml=fixture("ixbrl_micro.html", "rb")), keywords=DEFAULTS)
    assert res.status == "NOT DISCLOSED" and res.keywords_found == []
    assert t.keyword_cell(res) == "none found"


def test_no_document_means_keywords_not_searched():
    res = run(company("01234567", profile="profile_new.html"), keywords=DEFAULTS)
    assert res.keywords_found is None and t.keyword_cell(res) is None


def test_without_keywords_nothing_is_searched():
    res = run(company("01234567", filings=TAGGED, xhtml=fixture("ixbrl_tagged.html", "rb")))
    assert res.keywords_found is None


def test_ixbrl_already_searched_so_pdf_fallback_isnt_searched_again():
    full = filing_page(("01 Mar 2026", "AA", "Full accounts made up to 31 December 2025", True, True))
    res = run(company("01234567", filings=full, xhtml=fixture("ixbrl_micro.html", "rb"),
                      pdf=fixture("text_small_company.pdf", "rb")), keywords=["gross profit"])
    assert res.turnover == 5_432_100 and res.keywords_found == []      # the iXBRL was the document searched


@needs_ocr
def test_page_limit_is_reported():
    res = run(company("01234567", filings=SCANNED, pdf=fixture("scanned_annual_report.pdf", "rb")),
              keywords=["hedging"], max_pages=2)
    assert "keywords searched in 2 of 3 pages" in res.note


def test_turned_off_keyword_ocr_is_reported():
    res = run(company("01234567", filings=SCANNED, pdf=fixture("scanned_annual_report.pdf", "rb")),
              keywords=["hedging"], read_scanned=False)
    assert "keywords not searched in scanned pages" in res.note


@needs_ocr
def test_keyword_ocr_switch_off_still_finds_turnover():
    res = run(company("01234567", filings=SCANNED, pdf=fixture("scanned_annual_report.pdf", "rb")),
              keywords=["joint ventures"], ocr_keywords=False)
    assert res.status == "VERIFY" and res.keywords_found == []
    assert "keyword search of scanned PDFs is turned off" in res.note


# ---------------------------------------------------------------- spreadsheet output


def sheet(rows):
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return read_upload(buf.getvalue(), "c.xlsx")[1].active


def test_keywords_columns_and_summary():
    src = sheet([["Company No"], ["00445790"], ["16900000"]])
    routes = company("00445790", filings=TAGGED, xhtml=fixture("ixbrl_tagged.html", "rb"))
    routes.update(company("16900000", profile="profile_new.html"))
    out_wb, out = t.new_output()
    counts = t.process_sheet(src, {"name": None, "number": 1}, out, web=FakeWeb(routes), keywords=DEFAULTS)
    heads = [c.value for c in out[1]]
    count_col, found_col = heads.index("Keyword Count") + 1, heads.index("Keywords Found") + 1
    assert out.cell(2, count_col).value == 3          # the filing says "financial risks", not "financial risk"
    assert out.cell(2, found_col).value == "derivative (1), presentation currency (1), " \
                                           "derivative financial instruments (1)"
    assert out.cell(3, count_col).value is None and out.cell(3, found_col).value is None   # nothing to search
    assert counts["_with_keywords"] == 1

    t.add_summary(out_wb, counts, "c.xlsx", DEFAULTS)
    summary = out_wb["Summary"]
    cells = [c.value for row in summary.iter_rows() for c in row if c.value is not None]
    assert "Keywords" in cells and any(str(v).startswith("Searched for 162: foreign exchange") for v in cells)
    assert summary.cell(12, 1).value == "Total" and summary.cell(12, 2).value == 2   # statuses only


def test_no_keywords_leaves_keyword_columns_blank():
    src = sheet([["Company No"], ["00445790"]])
    routes = company("00445790", filings=TAGGED, xhtml=fixture("ixbrl_tagged.html", "rb"))
    _wb, out = t.new_output()
    t.process_sheet(src, {"name": None, "number": 1}, out, web=FakeWeb(routes))
    heads = [c.value for c in out[1]]
    assert out.cell(2, heads.index("Keyword Count") + 1).value is None
    assert out.cell(2, heads.index("Keywords Found") + 1).value is None


# ---------------------------------------------------------------- time limit


class FakeClock:
    """time.monotonic() that moves on 25 seconds every time it's read."""

    def __init__(self):
        self.now = 0.0

    def __call__(self):
        self.now += 25
        return self.now


def test_deadline_already_passed_reads_nothing():
    import time
    got = t.read_pdf(fixture("text_small_company.pdf", "rb"), deadline=time.monotonic() - 1)
    assert got.timed_out and got.pages_checked == 0 and got.turnover is None


def test_no_time_limit_and_no_page_limit_by_default():
    import inspect
    assert inspect.signature(t.check).parameters["time_limit"].default == 0
    assert inspect.signature(t.read_pdf).parameters["max_pages"].default is None
    assert inspect.signature(t.process_sheet).parameters["max_pages"].default is None


@needs_ocr
def test_time_limit_stops_reading_and_says_so(monkeypatch):
    monkeypatch.setattr(t.time, "monotonic", FakeClock())        # limit hit after 2 pages
    res = run(company("01234567", filings=SCANNED, pdf=fixture("scanned_annual_report.pdf", "rb")),
              keywords=["joint ventures", "hedging"], time_limit=1)
    assert res.status == "CHECK PDF" and res.turnover is None
    assert "stopped at the 1-minute time limit after 2 of 3 pages" in res.note
    assert "keywords searched in 2 of 3 pages (1-minute time limit)" in res.note
    assert res.keywords_found == ["joint ventures"]               # what was read still counts


@needs_ocr
def test_generous_time_limit_changes_nothing():
    res = run(company("01234567", filings=SCANNED, pdf=fixture("scanned_annual_report.pdf", "rb")),
              keywords=["joint ventures"], time_limit=30)
    assert res.status == "VERIFY" and res.turnover == 1_234_500_000 and "time limit" not in res.note


def test_singular_isnt_found_inside_plural_in_clean_text():
    k = t.Keywords(["forward contract", "forward contracts"])
    k.scan("the company enters into forward contracts where appropriate")
    assert k.ranked() == {"forward contracts": 1}


def test_counts_are_whole_words_not_parts_of_longer_ones():
    k = t.Keywords(["derivative", "derivatives", "forward contract", "forward contracts"])
    k.scan("A derivative. Two derivatives. One forward contract and three forward contracts.")
    assert k.ranked() == {"derivative": 1, "derivatives": 1, "forward contract": 1, "forward contracts": 1}
