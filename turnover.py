#!/usr/bin/env python3
"""
Automation Tool: for each company in a spreadsheet, find its latest annual accounts on
Companies House and pull out turnover (or revenue), prior-year turnover, cash and debtors,
plus the persons with significant control and, optionally, which keywords the accounts mention.

How it reads the accounts, best first:
  1. iXBRL (electronically filed accounts): the turnover figure is tagged, read exactly.
  2. iXBRL text, when the figure is shown but not tagged.
  3. PDF: text PDFs are read directly; scanned page images are read with OCR and marked
     VERIFY, because OCR can misread digits.

Small and micro companies are allowed to leave the profit and loss account out of what
they file, so for many companies there is no turnover to find. Those rows are flagged
with the reason instead.

Usage:
    python turnover.py clients.xlsx --name-col "Company Name" --number-col "Company No"

Install:
    pip install -r requirements.txt
"""

import argparse
import re
import sys
import threading
import time
import traceback
from urllib.parse import quote_plus
from dataclasses import dataclass, field
from datetime import date, datetime

import numpy as np
import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rapidfuzz import fuzz

BASE = "https://find-and-update.company-information.service.gov.uk"
WEB_DELAY = 1.0           # seconds between requests for the whole app, shared by every user
MAX_DELAY = 8.0           # the pace slows to at most this after Companies House pushes back
COOL_DOWN = 120           # seconds everyone pauses after a refusal (HTTP 403) or rate limit (429)
RETRIES = 4               # attempts per request for network errors and server errors
MATCH_THRESHOLD = 90      # name similarity needed to accept a search result
FILING_PAGES = 20         # pages of filing history (25 filings each) to look through
STALE_DAYS = 730          # accounts older than this get a note

# iXBRL tags for each figure, in order of preference (FRS 102, IFRS, old UK GAAP names)
TURNOVER_TAGS = ["TurnoverRevenue", "Revenue", "TurnoverGrossOperatingRevenue"]
CASH_TAGS = ["CashBankOnHand", "CashBankInHand", "CashAndCashEquivalents", "CashCashEquivalents"]
DEBTOR_TAGS = ["Debtors", "TradeAndOtherCurrentReceivables", "CurrentTradeAndOtherReceivables",
               "TradeAndOtherReceivables"]
ACCOUNT_TYPES = {"AA", "AAMD"}  # accounts, amended accounts
STATUSES = ["FOUND", "VERIFY", "NOT DISCLOSED", "CHECK PDF", "NOT FOUND", "ERROR"]

# ---------------------------------------------------------------- results


@dataclass
class Result:
    status: str = "NOT FOUND"       # one of STATUSES
    turnover: float | None = None
    prior_turnover: float | None = None
    cash: float | None = None
    debtors: float | None = None
    currency: str = ""
    period_end: str = ""
    company_name: str = ""
    company_number: str = ""
    company_status: str = ""
    accounts_type: str = ""
    filed_on: str = ""
    method: str = ""
    note: str = ""
    pdf_link: str = ""
    last_accounts: str | None = None    # from the company page: date, '' = none filed, None = unknown
    pscs: list | None = None            # [(name, is_corporate)] of active PSCs; None = not looked up
    keywords_found: list | None = None  # None = keywords weren't searched (no document, or none given)
    keyword_counts: dict = field(default_factory=dict)  # keyword -> times it appears
    notes: list = field(default_factory=list)

    @property
    def company_link(self):
        return f"{BASE}/company/{self.company_number}" if self.company_number else ""


class SourceError(Exception):
    """Companies House couldn't be reached or refused the request."""


# ---------------------------------------------------------------- web access


class Pace:
    """One pace for every request the app makes, whichever user it's for. All the app's
    requests come from the same server address, and the website blocks that address
    (HTTP 403) when the total gets too high, so users take turns at a steady rate. When
    Companies House pushes back, everyone pauses and the pace slows; it speeds back up
    gradually while requests go through."""

    def __init__(self, delay=WEB_DELAY):
        self.base = self.delay = delay
        self.lock = threading.Lock()
        self.next_at = 0.0

    def wait_turn(self):
        with self.lock:                          # one request at a time, in order
            pause = self.next_at - time.monotonic()
            if pause > 0:
                time.sleep(pause)
            self.next_at = time.monotonic() + self.delay

    def pushed_back(self, cool_down):
        with self.lock:
            self.delay = min(max(self.delay * 2, 2.0), MAX_DELAY)
            self.next_at = max(self.next_at, time.monotonic() + cool_down)

    def went_through(self):
        self.delay = max(self.base, self.delay * 0.98)


PACE = Pace()                                    # shared by every Web() in this process


class Web:
    def __init__(self, delay=None, pace=None):
        self.s = requests.Session()
        self.s.headers["User-Agent"] = "Mozilla/5.0 (Automation Tool)"
        self.pace = pace or (Pace(delay) if delay is not None else PACE)

    def get(self, path, params=None):
        """GET a page; None for 404. Retries network errors, server errors, rate limits and
        temporary refusals, pausing the whole app's requests when Companies House pushes back."""
        url = path if path.startswith("http") else BASE + path
        problem = ""
        for attempt in range(RETRIES):
            self.pace.wait_turn()
            try:
                r = self.s.get(url, params=params, timeout=60)
            except requests.RequestException as e:
                problem = f"network error ({type(e).__name__})"
                time.sleep(5 * (attempt + 1))
                continue
            if r.status_code == 404:
                self.pace.went_through()
                return None
            if r.status_code in (403, 429):
                problem = ("Companies House is refusing requests from this server (HTTP 403), usually a "
                           "temporary block after heavy use" if r.status_code == 403
                           else "rate limited by Companies House")
                self.pace.pushed_back(COOL_DOWN * (attempt + 1))
                continue
            if r.status_code >= 500:
                problem = f"Companies House server error (HTTP {r.status_code})"
                time.sleep(10 * (attempt + 1))
                continue
            if r.status_code >= 400:
                raise SourceError(f"Companies House refused the request (HTTP {r.status_code})")
            self.pace.went_through()
            return r
        raise SourceError(f"{problem}, tried {RETRIES} times; run this row again later")

    def soup(self, path, params=None):
        r = self.get(path, params)
        return BeautifulSoup(r.text, "html.parser") if r is not None else None


# ---------------------------------------------------------------- finding the company


def normalise(name):
    s = str(name or "").upper().replace("&", " AND ")
    s = re.sub(r"\bP\.\s*L\.\s*C\.?", "PLC", s)              # "P.L.C." as the register writes it
    s = re.sub(r"\bL\.\s*T\.\s*D\.?", "LTD", s)
    s = re.sub(r"[^\w\s]", " ", s)
    s = re.sub(r"\bLTD\b", "LIMITED", s)
    s = re.sub(r"\bPLC\b", "PUBLIC LIMITED COMPANY", s)
    s = re.sub(r"\b(PUBLIC LIMITED COMPANY|LIMITED LIABILITY PARTNERSHIP|LIMITED|COMPANY)\b", " ", s)
    s = re.sub(r"^THE\s+", "", s.strip())
    return re.sub(r"\s+", " ", s).strip()


def name_score(a, b):
    na, nb = normalise(a), normalise(b)
    return 100.0 if na == nb else max(fuzz.token_sort_ratio(na, nb), fuzz.ratio(na, nb))


def clean_number(value):
    """Company numbers as people type them, or as Excel mangles them: 445790, 445790.0,
    'sc 5261', ' 00445790 ' -> '00445790', 'SC005261'. '' when there's nothing usable."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = re.sub(r"[^A-Za-z0-9]", "", str(value)).upper()
    if s.isdigit():
        return s.zfill(8)
    m = re.fullmatch(r"([A-Z]{2})(\d{1,6})", s)
    return m.group(1) + m.group(2).zfill(6) if m else s


def read_profile(soup, res):
    """Name, status and accounts facts from a company page. Returns the 'last accounts' date
    text, '' when the company hasn't filed accounts yet, or None when the page doesn't say."""
    h1 = soup.select_one("#company-name") or soup.find("h1")
    res.company_name = h1.get_text(" ", strip=True) if h1 else res.company_name
    st = soup.select_one("#company-status")
    res.company_status = st.get_text(" ", strip=True) if st else ""
    if res.company_status.lower() in ("dissolved", "liquidation", "administration", "receivership",
                                      "insolvency proceedings", "voluntary arrangement", "closed"):
        res.notes.append(f"company status: {res.company_status.lower()}")
    text = re.sub(r"\s+", " ", soup.get_text(" ", strip=True))
    if re.search(r"Accounts overdue", text, re.I):
        res.notes.append("accounts overdue at Companies House")
    m = re.search(r"Last accounts made up to (\d{1,2} \w+ \d{4})", text)
    if m:
        return m.group(1)
    if re.search(r"First accounts made up to", text):
        return ""
    return None


def find_company(web, name, number, res):
    """Look the company up by number (exact) or name (search). Fills in res; returns the
    company page soup, or None with res.note saying why."""
    number = clean_number(number)
    if number:
        soup = web.soup(f"/company/{number}")
        if not soup:
            res.note = f"company number {number} not on the register"
            return None
        res.company_number = number
        res.last_accounts = read_profile(soup, res)
        if name and res.company_name and name_score(name, res.company_name) < 70:
            res.notes.append(f"the number belongs to {res.company_name}, check it matches '{name}'")
        return soup

    if not name:
        res.note = "no company name or number in this row"
        return None
    search = web.soup("/search/companies", {"q": name})
    hits = []
    for li in (search.select("#results li") if search else []):
        a = li.find("a", href=re.compile(r"^/company/[A-Z0-9]+"))
        if a:
            hits.append((a.get_text(" ", strip=True), a["href"].split("/")[2]))
    if not hits:
        res.note = "company not found on the register"
        return None
    best = max(hits, key=lambda h: name_score(name, h[0]))
    score = name_score(name, best[0])
    if score < MATCH_THRESHOLD:
        res.note = f"no confident match (closest: {best[0]}, {score:.0f}%), add the company number"
        return None
    if any(n != best[1] and name_score(name, t) >= MATCH_THRESHOLD for t, n in hits):
        res.notes.append("several companies with near-identical names, add the company number to be sure")
    res.company_name, res.company_number = best
    soup = web.soup(f"/company/{best[1]}")
    if soup:
        res.last_accounts = read_profile(soup, res)
    return soup or BeautifulSoup("", "html.parser")


# ---------------------------------------------------------------- latest accounts filing


def latest_accounts(web, number, pages=FILING_PAGES):
    """Return (date, description, pdf_url, xhtml_url) of the newest annual accounts filing,
    or None. The website ignores its own 'accounts only' filter, so this pages through the
    whole filing history (newest first) until it meets an accounts filing."""
    for page in range(1, pages + 1):
        soup = web.soup(f"/company/{number}/filing-history", {"page": page})
        rows = soup.select("#fhTable tr")[1:] if soup else []
        for tr in rows:
            tds = tr.find_all("td")
            if len(tds) < 3 or tds[1].get_text(strip=True) not in ACCOUNT_TYPES:
                continue
            description = tds[2].get_text(" ", strip=True).split(" View PDF")[0]
            if re.match(r"\s*(interim|initial)\s+accounts", description, re.I):
                continue                    # not the annual accounts
            links = {a["href"] for a in tr.find_all("a", href=True)}
            pdf = next((l for l in links if "format=pdf" in l), "")
            xhtml = next((l for l in links if "format=xhtml" in l), "")
            return (tds[0].get_text(" ", strip=True), description,
                    BASE + pdf if pdf else "", BASE + xhtml if xhtml else "")
        if not rows or not soup.find("a", href=re.compile(rf"[?&]page={page + 1}\b")):
            return None
    return None


# ---------------------------------------------------------------- reading amounts from lines of text

NOTE_REF = r"(?:\(?\s*notes?\b[^)]*\)?)?"
LABEL = re.compile(r"^\s*(?:group\s*|total\s*|net\s*)?(?:turnover|revenue|sales)s?\b\s*(?:\((?!\s*\d)[^)]*\))?\s*"
                   + NOTE_REF + r"(.*)$", re.I)   # allows "Revenue (including ...)"
CASH_LABEL = re.compile(r"^\s*cash(?:\s*at\s*bank(?:\s*and\s*in\s*hand)?|\s*and\s*cash\s*equivalents|\s*in\s*hand"
                        r"|\s*and\s*short\s*-?\s*term\s*deposits|\s*and\s*bank\s*balances)?"
                        r"\b(?!\s*flow)\s*" + NOTE_REF + r"(.*)$", re.I)
DEBTORS_LABEL = re.compile(r"^\s*(?:total\s*)?(?:[a-z]+\s*and\s*other\s*)?(?:debtors|receivables)\b"
                           r"(?:\s*:?\s*amounts\s*falling\s*due\s*within\s*one\s*year)?\s*"
                           + NOTE_REF + r"(.*)$", re.I)       # "Trade (or Customer) and other receivables"
WHOLE = re.compile(r"\d{1,3}(?:,\d{3})+|\d{3,}")         # 1,234,567 or 1234567
OCR_DOTS = re.compile(r"\d{1,3}(?:\.\d{3})+")             # 72.886: OCR read the comma as a dot
DECIMAL = re.compile(r"\d{1,3}(?:,\d{3})*\.\d{1,2}")       # 2,151.2 (accounts in £m)
OCR_DOTS_DECIMAL = re.compile(r"\d{1,3}(?:\.\d{3})+\.\d{1,2}")  # 1.771.0: OCR read 1,771.0 with dots


def amounts(rest, in_thousands_or_millions=False):
    """Pull money amounts out of the text after a label, skipping note references
    (1-2 digit numbers like '3' or '2,3'). A small decimal like '2.1' could be a note number,
    so those only count when the row has several; 1,000.0 and up always count. When the
    figures are in £000 or £m a 1-2 digit number can be a real amount ('Cash 71 125'), so
    then only a small number straight after the label in a row of three or more is a note."""
    found, small = [], []
    for tok in re.split(r"\s+", rest.replace("£", "").strip()):
        if re.fullmatch(r"\(?(19|20)\d{2}:", tok):        # "(2024: £1,234)" names the prior year
            continue
        tok = tok.strip("|:;")
        if tok.lower() in ("0", "-", "–", "—", "nil"):         # zero this year (or last) holds its column
            found.append((0.0, False, False))
            continue
        neg = tok.startswith("(") or tok.startswith("-")
        core = tok.strip("()-–")
        if re.fullmatch(r"[0,.]+", core):                    # "£000" is a unit heading, not zero
            continue
        if WHOLE.fullmatch(core):
            found.append((float(core.replace(",", "")), neg, False))
        elif in_thousands_or_millions and re.fullmatch(r"\d{1,2}", core):
            small.append(len(found))
            found.append((float(core), neg, False))
        elif OCR_DOTS.fullmatch(core):
            found.append((float(core.replace(".", "")), neg, False))
        elif DECIMAL.fullmatch(core):
            v = float(core.replace(",", ""))
            found.append((v, neg, v < 1000))
        elif OCR_DOTS_DECIMAL.fullmatch(core):
            whole, _, dec = core.rpartition(".")
            found.append((float(whole.replace(".", "") + "." + dec), neg, False))
    if small == [0] and len(found) >= 3:
        found = found[1:]                                # a note number before the amounts
    if sum(1 for f in found if f[2]) < 2:
        found = [f for f in found if not f[2]]
    return [-v if neg else v for v, neg, _ in found]


def figure_from_lines(lines, label=LABEL, start=0):
    """lines: strings, one per row of a table. Finds the first row starting with label and
    returns (current, prior, row_units, row_index) or None. prior is only given when the row
    is plainly 'this year, last year': exactly two amounts within a factor of 10 of each
    other (wider rows have extra columns, like 'before adjusting items'). The search begins at
    row start; rows above it still count for the units."""
    for idx in range(start, len(lines)):
        line = lines[idx]
        m = label.match(line)
        if not m or re.match(r"\s*[A-Za-z]{2,}", m.group(1)):
            continue                                    # "Sales of £109,016 ..." is a sentence, not a row
        # Stop at the next word: either the label carries on ("Revenue from sale of goods", a
        # sub-line, so no amounts before it) or OCR merged in a neighbouring column's row.
        rest = re.split(r"(?i)(?!nil(?![a-z]))[a-z]{3,}", m.group(1), maxsplit=1)[0]
        units = unit_multiplier(rest)
        context = units if units[0] > 1 else (unit_multiplier(" | ".join(lines[max(0, idx - 8):idx])))
        if context[0] == 1:
            context = heading_units(lines, idx) or context
        vals = amounts(rest, context[0] > 1)
        if vals:
            prior = vals[1] if len(vals) == 2 and vals[0] and 0.1 <= abs(vals[1] / vals[0]) <= 10 else None
            return vals[0], prior, (units if units[0] > 1 else None), idx
    return None


def turnover_from_lines(lines):
    return figure_from_lines(lines, LABEL)


MONTHS = r"jan(uary)?|feb(ruary)?|mar(ch)?|apr(il)?|may|june?|july?|aug(ust)?|sep(tember)?|oct(ober)?|nov(ember)?|dec(ember)?"
HEADER_TOKEN = re.compile(r"(notes?|[£Ef€$]?[’']?(m|000|million)|[£€$]|\d{4}|\d{1,2}|restated|group|company|"
                          rf"weeks|ended|year|period|to|total|and|as|at|({MONTHS})\d{{0,4}})", re.I)


def heading_units(lines, idx, reach=40):
    """Units from a column-heading line further up ('2025 2024 £m £m'). Only lines made of
    nothing but heading words count, so a stray 'm' in a sentence can't change the units."""
    for line in reversed(lines[max(0, idx - reach):idx]):
        tokens = [t for t in re.split(r"[\s|]+", line) if t]
        if tokens and all(HEADER_TOKEN.fullmatch(t) for t in tokens):
            units = unit_multiplier(line)
            if units[0] > 1:
                return units
    return None


def scaled(found, lines):
    """Apply the units (£, £000, £m) to a figure_from_lines result. Units come from the row
    itself, the lines just above it, or a column-heading line further up; never the whole
    page, where a stray 'm' would turn pounds into millions."""
    idx = found[3]
    near = unit_multiplier(" | ".join(lines[max(0, idx - 8):idx]))
    mult, label = found[2] or (near if near[0] > 1 else None) or heading_units(lines, idx) or near
    return round(found[0] * mult), (round(found[1] * mult) if found[1] is not None else None), label


def balance_figure(lines, label):
    """Cash or debtors from a balance sheet's rows. Balance sheets list non-current receivables
    before current assets, so look below the 'Current assets' heading first."""
    for i, line in enumerate(lines):
        if re.match(r"\s*current\s*assets\b(?!\s*[\d(£-])", line, re.I):   # the heading, not a total
            found = figure_from_lines(lines, label, start=i)
            # Not readable under the heading: better blank than a non-current figure from above.
            return scaled(found, lines) if found else None
    found = figure_from_lines(lines, label)
    return scaled(found, lines) if found else None


def unit_multiplier(text):
    """Column headers like '£m', '€ million' or '£000' say what the figures are in.
    OCR often reads £ as E and sometimes drops the symbol altogether."""
    tokens = set(re.split(r"[\s|]+", text))
    if any(re.fullmatch(r"[£Ef€$]?[’']?(m|million)", t) for t in tokens):   # £m, £'m, Em, f'm, € million
        return 1_000_000, "m"
    if any(re.fullmatch(r"[£E€$]?[’']?000|[£€$]k|[£E€$]?[’']000s?", t) for t in tokens):
        return 1_000, "000"
    return 1, ""


def page_currency(text):
    """GBP unless the page shows another currency. '' when no symbol is visible at all."""
    if re.search(r"€|\beuros?\b", text, re.I):
        return "EUR"
    if re.search(r"US\$|\$m|\$'?000|\bUSD\b|dollars", text):
        return "USD"
    if re.search(r"£|\b[Ef][m’']\b|\bE\s*000|\bGBP\b|sterling", text):
        return "GBP"
    return ""


# ---------------------------------------------------------------- keywords


def _norm(text):
    """Lowercase, rejoin words split across lines ('ex-\nchange'), punctuation to spaces."""
    text = re.sub(r"-\s*\n\s*", "", str(text).lower())
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", text)).strip()


def parse_keywords(text):
    """Keywords typed one per line, or separated by commas or semicolons."""
    return [k.strip() for k in re.split(r"[\n,;]+", text or "") if k.strip()]


class Keywords:
    """Finds which of a list of keywords appear in a document's text, and how often.

    Matching ignores case and punctuation, and counts whole words or phrases only, so
    'derivative' doesn't count 'derivatives' and 'FX' doesn't match inside other words. In OCR
    text, a phrase or long word (8+ letters) that isn't found that way is looked for with spaces
    ignored, because OCR often runs words together ("Foreignexchange")."""

    def __init__(self, words):
        self.words, seen = [], set()
        for w in words:
            n = _norm(w)
            if n and n not in seen:
                seen.add(n)
                self.words.append(w.strip())
        self._tests = []
        for w in self.words:
            n = _norm(w)
            compact = n.replace(" ", "")
            whole = re.compile(rf"(?<![a-z0-9]){re.escape(n)}(?![a-z0-9])")
            self._tests.append((w, whole, compact if " " in n or len(compact) >= 8 else None))
        self.found = set()
        self.counts = {}

    def scan(self, text, ocr=False):
        """Add one piece of text (a page, or a whole document) to the counts. ocr=True for text
        read by OCR, which allows the spaces-ignored fallback."""
        if not text:
            return
        t = _norm(text)
        tc = t.replace(" ", "")
        for w, whole, compact in self._tests:
            n = len(whole.findall(t))
            if not n and compact and ocr:               # OCR may have run the words together
                n = tc.count(compact)
            if n:
                self.found.add(w)
                self.counts[w] = self.counts.get(w, 0) + n

    @property
    def done(self):
        return len(self.found) == len(self.words)

    def result(self):
        """Found keywords in the order they were given."""
        return [w for w in self.words if w in self.found]

    def ranked(self):
        """{keyword: count} for the keywords found, most mentioned first (ties keep the given order)."""
        order = {w: i for i, w in enumerate(self.words)}
        return dict(sorted(self.counts.items(), key=lambda kv: (-kv[1], order[kv[0]])))


# ---------------------------------------------------------------- iXBRL


def parse_number(el):
    txt = el.get_text("", strip=True)
    fmt = (el.get("format") or "").lower()
    if re.fullmatch(r"[-–—]*", txt) or "zerodash" in fmt or "fixed-zero" in fmt:
        return 0.0
    if "numcommadecimal" in fmt or "num-comma-decimal" in fmt:
        txt = txt.replace(".", "").replace(" ", "").replace(",", ".")
    else:
        txt = txt.replace(",", "").replace(" ", "")
    try:
        v = float(re.sub(r"[^\d.]", "", txt))
    except ValueError:
        return None
    v *= 10 ** int(el.get("scale") or 0)
    return -v if el.get("sign") == "-" else v


def months_between(start, end):
    try:
        a, b = date.fromisoformat(start), date.fromisoformat(end)
    except (TypeError, ValueError):
        return None
    return round((b - a).days / 30.44)


def _ix_facts(soup):
    """{tag name: [(rank, end, start, value, unit)]} for every number in an iXBRL document,
    worked out once per document. rank 0 = no dimensions, 1 = the consolidated group figure,
    2 = the amount due within one year (how many filings tag current debtors); other
    breakdowns (a segment, a subsidiary, an age band...) aren't headline figures."""
    cached = soup.__dict__.get("_ix_facts")
    if cached is not None:
        return cached
    contexts = {}
    for c in soup.find_all(re.compile(r"(^|:)context$")):
        end = c.find(re.compile(r"(^|:)(enddate|instant)$"))
        start = c.find(re.compile(r"(^|:)startdate$"))
        members = [m.get_text(strip=True) for m in c.find_all(re.compile(r"(^|:)explicitmember$"))]
        dims = c.find(re.compile(r"(^|:)(segment|scenario)$"))
        if not dims:
            rank = 0
        elif members and all(m.split(":")[-1].startswith("Consolidated") for m in members):
            rank = 1
        elif members and all(re.search(r"^Consolidated|(?<!Non)CurrentFinancialInstruments$|WithinOneYear$",
                                       m.split(":")[-1]) for m in members):
            rank = 2
        else:
            rank = None
        contexts[c.get("id")] = (end.get_text(strip=True) if end else "",
                                 start.get_text(strip=True) if start else "", rank)
    units = {}
    for u in soup.find_all(re.compile(r"(^|:)unit$")):
        m = u.find(re.compile(r"(^|:)measure$"))
        units[u.get("id")] = m.get_text(strip=True).split(":")[-1] if m else ""
    facts = {}
    for el in soup.find_all(re.compile(r"(^|:)nonfraction$")):
        end, start, rank = contexts.get(el.get("contextref"), ("", "", None))
        v = parse_number(el)
        if v is not None and rank is not None and end:
            facts.setdefault(el.get("name", "").split(":")[-1], []).append(
                (rank, end, start, v, units.get(el.get("unitref"), "")))
    soup.__dict__["_ix_facts"] = facts
    return facts


def ix_value(soup, tags):
    """(current, prior, unit, end, start) for the first of tags the document reports, or None."""
    facts = _ix_facts(soup)
    for tag in tags:
        if not facts.get(tag):
            continue
        by_end = {}
        for _rank, end, start, v, unit in sorted(facts[tag], key=lambda f: (f[0], f[1])):
            by_end.setdefault(end, (v, unit, start))
        ends = sorted(by_end, reverse=True)
        v, unit, start = by_end[ends[0]]
        return v, (by_end[ends[1]][0] if len(ends) > 1 else None), unit, ends[0], start
    return None


TOP = re.compile(r"top:\s*(-?[\d.]+)", re.I)
LEFT = re.compile(r"left:\s*(-?[\d.]+)", re.I)


def ix_lines(soup):
    """The document's rows of text: table rows, plus rows rebuilt from positioned text blocks.
    Filings converted from PDF place every word and number by its coordinates, so blocks at
    the same height on the same page are one row, read left to right."""
    cached = soup.__dict__.get("_ix_lines")
    if cached is not None:
        return cached
    lines = [" ".join(td.get_text(" ", strip=True) for td in tr.find_all(["td", "th"]))
             for tr in soup.find_all("tr")]
    pages = {}
    for el in soup.find_all(style=TOP):
        if el.find(style=TOP):
            continue                                    # only the innermost blocks hold the text
        text = el.get_text(" ", strip=True)
        if not text:
            continue
        top, left = TOP.search(el["style"]), LEFT.search(el["style"])
        rows = pages.setdefault(id(el.parent), {})
        rows.setdefault(round(float(top.group(1)), 1), []).append((float(left.group(1)) if left else 0.0, text))
    for rows in pages.values():
        for top in sorted(rows):
            lines.append(" ".join(t for _x, t in sorted(rows[top])))
    soup.__dict__["_ix_lines"] = lines
    return lines


def read_ixbrl(html):
    """Return (current, prior, currency, period_end, method, months) for turnover, or None.
    months is the length of the accounting period when the filing says. html may be a soup."""
    soup = html if isinstance(html, BeautifulSoup) else BeautifulSoup(html, "html.parser")
    tagged = ix_value(soup, TURNOVER_TAGS)
    if tagged:
        v, prior, unit, end, start = tagged
        return v, prior, (unit or "GBP").upper(), end, "iXBRL tagged figure", months_between(start, end)
    lines = ix_lines(soup)                              # shown but not tagged: read the rows as text
    found = statement_figure(lines, LABEL, STATEMENT_HEAD)
    if found:
        current, prior, _ = found
        return current, prior, "GBP", "", "iXBRL table (untagged)", None
    return None


def statement_figure(lines, label, heading, span=60):
    """A figure from the rows just below a statement heading (the income statement for
    turnover, the balance sheet for cash and debtors), else from anywhere in the document.
    Returns scaled (current, prior, unit) or None."""
    starts = [i for i, l in enumerate(lines) if heading.search(l) and not NOTES_HEAD.search(l)]
    pick = balance_figure if heading is BALANCE_HEAD else (lambda w, lb: scaled(f, w) if (f := figure_from_lines(w, lb)) else None)
    for i in starts + [None]:
        found = pick(lines if i is None else lines[i:i + span], label)
        if found:
            return found
    return None


def read_ixbrl_balance(soup):
    """(cash, debtors) at the balance sheet date: tagged figures, else read from the rows."""
    out = []
    for tags, label in ((CASH_TAGS, CASH_LABEL), (DEBTOR_TAGS, DEBTORS_LABEL)):
        tagged = ix_value(soup, tags)
        if tagged:
            out.append(tagged[0])
            continue
        found = statement_figure(ix_lines(soup), label, BALANCE_HEAD)
        out.append(found[0] if found else None)
    return tuple(out)


def ixbrl_text(soup):
    """The readable text of an iXBRL document, without the hidden tagging header."""
    _ix_facts(soup)                                     # read the tags before the header goes
    for header in soup.find_all(re.compile(r"(^|:)header$")):
        header.decompose()
    return soup.get_text(" ", strip=True)


# ---------------------------------------------------------------- PDF + OCR

_ocr = None
OCR_SLOTS = threading.BoundedSemaphore(1)   # scanned pages are read one at a time across all users


def ocr_turn(progress=None):
    """Wait for the scanned-PDF reader. Reading uses a lot of memory and CPU, so the app reads
    one page at a time; two users' reports take turns page by page."""
    if not OCR_SLOTS.acquire(blocking=False):
        if progress:
            progress("waiting for the scanned-PDF reader (another user's report is being read)")
        OCR_SLOTS.acquire()


class OcrUnavailable(Exception):
    pass


_ocr_problem = None
_ocr_start = threading.Lock()


def ocr_available():
    """Whether scanned PDFs can be read here: the OCR package is installed and actually starts.
    (It only installs on Python 3.12 or earlier, and on Linux it needs the system libraries in
    packages.txt.)"""
    try:
        ocr_engine()
        return True
    except OcrUnavailable:
        return False


def ocr_engine():
    """The OCR reader, started once and shared. If it can't start, the reason goes to the log
    (Manage app > logs on Streamlit Cloud) and every later call fails fast with it."""
    global _ocr, _ocr_problem
    if _ocr is None:
        with _ocr_start:
            if _ocr is None:
                if _ocr_problem:
                    raise OcrUnavailable(_ocr_problem)
                try:
                    from rapidocr_onnxruntime import RapidOCR
                    _ocr = RapidOCR()
                except Exception as e:          # missing package, model or system library
                    _ocr_problem = f"{type(e).__name__}: {e}"
                    print(f"Scanned-PDF reader couldn't start: {_ocr_problem}", file=sys.stderr)
                    traceback.print_exc(file=sys.stderr)
                    raise OcrUnavailable(_ocr_problem) from e
    return _ocr


def _rows(boxes):
    """Group OCR boxes [(y_mid, x_left, height, payload)] into rows, each sorted left to right."""
    if not boxes:
        return []
    tol = 0.6 * sorted(b[2] for b in boxes)[len(boxes) // 2]
    rows, current = [], []
    for b in sorted(boxes, key=lambda b: b[0]):
        if current and b[0] - current[-1][0] > tol:
            rows.append(current)
            current = []
        current.append(b)
    rows.append(current)
    return [sorted(r, key=lambda b: b[1]) for r in rows]


def _box(pts, payload):
    xs, ys = [p[0] for p in pts], [p[1] for p in pts]
    return ((min(ys) + max(ys)) / 2, min(xs), max(ys) - min(ys), payload)


def ocr_rows(image):
    """OCR a whole page; returns one string per row of text, left to right."""
    res, _ = ocr_engine()(np.asarray(image.convert("RGB")), use_cls=False)
    rows = _rows([_box(pts, text) for pts, text, _conf in res or []])
    return [" ".join(b[3] for b in r) for r in rows]


def ocr_skim(image):
    """Cheap pass: find all the text, but only read the first box on each row. That's enough to
    spot headings and row labels like 'Revenue', at a fraction of the cost of a full read."""
    arr = np.asarray(image.convert("RGB"))
    eng = ocr_engine()
    found, _ = eng(arr, use_cls=False, use_rec=False)
    if not found:
        return []
    firsts = [r[0] for r in _rows([_box(pts, pts) for pts in found])]
    crops = []
    for _y, _x, _h, pts in firsts:
        xs, ys = [int(p[0]) for p in pts], [int(p[1]) for p in pts]
        crops.append(arr[max(min(ys), 0):max(ys) + 1, max(min(xs), 0):max(xs) + 1])
    texts, _ = eng.text_rec(crops)
    return [t[0] for t in texts]


SKIM_LABEL = re.compile(r"^(group\s*|total\s*|net\s*)?(turnover|revenue|sales)s?\b", re.I)


def page_order(n):
    """Long annual reports put the financial statements after the strategic report and
    governance sections, so start a quarter of the way in and wrap round."""
    if n <= 40:
        return list(range(n))
    start = int(n * 0.25)
    return list(range(start, n)) + list(range(start))


STATEMENT_HEAD = re.compile(r"income\s*statement|profit\s*(and|&)\s*loss|statement\s*of\s*(comprehensive\s*)?income|"
                            r"statement\s*of\s*profit\s*(or|and)\s*loss", re.I)
NOTES_HEAD = re.compile(r"notes\s*to\s*the|notes\s*\(continued\)", re.I)


def is_statement_page(top_rows):
    """The page is a primary statement when its heading (first few rows) names one, and it
    isn't a note to the accounts (notes repeat 'Revenue' for subsidiaries and segments)."""
    top = " | ".join(top_rows[:8])
    return bool(STATEMENT_HEAD.search(top)) and not NOTES_HEAD.search(top)


BALANCE_HEAD = re.compile(r"balance\s*sheet|statement\s*of\s*financial\s*position", re.I)
BALANCE_SKIM = re.compile(r"^(cash|debtors|[a-z]+\s*and\s*other\s*receivables)", re.I)
BALANCE_PAGES_AFTER_TURNOVER = 15   # the balance sheet is close behind the income statement


def is_balance_page(top_rows):
    top = " | ".join(top_rows[:8])
    return bool(BALANCE_HEAD.search(top)) and not NOTES_HEAD.search(top)


@dataclass
class PdfResult:
    turnover: float | None = None
    prior: float | None = None
    unit: str = ""
    page: int = 0
    method: str = ""
    currency: str = ""
    cash: float | None = None
    debtors: float | None = None
    balance_method: str = ""
    pages: int = 0                # pages in the document
    pages_checked: int = 0
    keyword_pages: int = 0        # pages whose full text was searched for keywords
    scanned_skipped: bool = False
    keyword_ocr_skipped: bool = False
    timed_out: bool = False
    problem: str = ""


def read_pdf(data, max_pages=None, progress=None, read_scanned=True, keywords=None, ocr_keywords=True,
             deadline=None):
    """Find turnover on the income statement page and cash and debtors on the balance sheet of
    an accounts PDF. If keywords is given, every page is also searched for them; scanned pages
    are then read in full, which is slow (ocr_keywords=False searches only pages with real
    text). max_pages (None = all) and deadline (a time.monotonic() value, None = no limit) stop
    the reading early."""
    import pypdfium2 as pdfium
    out = PdfResult()
    try:
        pdf = pdfium.PdfDocument(data)
        out.pages = len(pdf)
    except Exception as e:
        out.problem = f"the accounts PDF couldn't be opened ({type(e).__name__})"
        return out
    order = page_order(out.pages)[:max_pages]
    since_turnover = 0
    for done, i in enumerate(order, 1):
        want_keywords = keywords is not None             # counts need every page, so never stop early
        need_turnover = out.turnover is None
        need_balance = out.cash is None or out.debtors is None
        if not need_turnover:
            since_turnover += 1
        if not want_keywords and not need_turnover and (
                not need_balance or since_turnover > BALANCE_PAGES_AFTER_TURNOVER):
            break
        if deadline is not None and time.monotonic() > deadline:
            out.timed_out = True
            break
        out.pages_checked = done
        try:
            page = pdf[i]
            text = page.get_textpage().get_text_range()
        except Exception:
            continue                                       # one unreadable page, try the rest
        if text.strip():                                   # PDF has real text, no OCR needed
            rows, method = text.splitlines(), "PDF text"
            if want_keywords:
                keywords.scan(text)
                out.keyword_pages += 1
        elif not read_scanned:
            out.scanned_skipped = True
            continue
        else:
            if progress:
                what = "searching scanned accounts" if want_keywords else "reading scanned accounts"
                progress(f"{what}, page {done} of {len(order)}")
            if want_keywords and not ocr_keywords:
                want_keywords, out.keyword_ocr_skipped = False, True
            ocr_turn(progress)
            try:
                image = page.render(scale=150 / 72).to_pil()
                if want_keywords:                          # need every word on the page
                    rows = ocr_rows(image)
                    keywords.scan("\n".join(rows), ocr=True)
                    out.keyword_pages += 1
                else:                                      # only the income statement and balance sheet
                    skim = ocr_skim(image)
                    wanted = (need_turnover and is_statement_page(skim)
                              and any(SKIM_LABEL.match(t.strip()) for t in skim)) or \
                             (need_balance and is_balance_page(skim)
                              and any(BALANCE_SKIM.match(t.strip()) for t in skim))
                    if not wanted:
                        continue
                    rows = ocr_rows(image)
                method = "PDF scan (OCR)"
            except OcrUnavailable:
                out.problem = ("this is a scanned PDF and the scanned-PDF reader isn't working on this "
                               "server, open the PDF to check")
                return out
            finally:
                OCR_SLOTS.release()
        heading = [r for r in rows if r.strip()]
        if need_turnover and is_statement_page(heading):
            found = turnover_from_lines(rows)
            if found:
                out.turnover, out.prior, out.unit = scaled(found, rows)
                out.page, out.method = i + 1, method
                out.currency = page_currency(" | ".join(rows))
        if need_balance and is_balance_page(heading):
            for attr, label in (("cash", CASH_LABEL), ("debtors", DEBTORS_LABEL)):
                if getattr(out, attr) is None:
                    found = balance_figure(rows, label)
                    if found:
                        setattr(out, attr, found[0])
                        out.balance_method = method
    return out


# ---------------------------------------------------------------- one company


def why_missing(description):
    """Reason a filing has no turnover, from its description; '' when the accounts type
    should normally include it (full, group, audited...)."""
    d = description.lower()
    if "dormant" in d:
        return "dormant company, no turnover"
    if "micro" in d:
        return "micro-entity accounts, turnover not required to be filed"
    if any(k in d for k in ("total exemption", "abridged", "small", "filleted")):
        return "small company accounts, profit and loss not filed"
    if "medium" in d:
        return "medium company accounts, which may leave turnover out"
    return ""


def read_pscs(web, number):
    """Active persons with significant control as [(name, is_corporate)]; [] when none are listed."""
    soup = web.soup(f"/company/{number}/persons-with-significant-control")
    out = []
    for el in (soup.select('[id^="psc-name-"]') if soup else []):
        idx = el["id"].rsplit("-", 1)[-1]
        status = soup.select_one(f"#psc-status-tag-{idx}")
        block = el.find_parent(class_=re.compile(r"^appointment-\d+$"))
        text = block.get_text(" ", strip=True) if block else ""
        if (status and "ceased" in status.get_text(" ", strip=True).lower()) or "Ceased on" in text:
            continue
        name = re.sub(r"\s+", " ", el.get_text(" ", strip=True))
        out.append((name, bool(re.search(r"Legal form|Registration number|Place registered", text))))
    return out


def apply_pdf(res, pdf):
    res.turnover, res.prior_turnover, res.method = pdf.turnover, pdf.prior, pdf.method
    # UK accounts often print no currency symbol at all ("'000", or just figures), so pounds
    # unless the page shows another currency (€, $, "euro", "US dollars").
    res.currency = pdf.currency or "GBP"
    res.method += f", page {pdf.page}" + (f", figures in {res.currency}{pdf.unit}" if pdf.unit else "")
    res.status = "VERIFY" if "OCR" in res.method else "FOUND"
    if res.status == "VERIFY":
        res.notes.append(f"read by OCR from page {pdf.page}, check the figures against the PDF")


def check(row, web, max_pages=None, progress=None, read_scanned=True, keywords=None,
          ocr_keywords=True, time_limit=0):
    """Look up one row: {'name': ..., 'number': ...}. keywords, if given, is a list of words or
    phrases to look for in the same accounts. time_limit is in minutes per company (0 = no
    limit); when it runs out, PDF reading stops and the row says so."""
    deadline = time.monotonic() + time_limit * 60 if time_limit else None
    res = Result()
    kw = Keywords(keywords) if keywords else None
    name = str(row.get("name") or "").strip()
    profile = find_company(web, name, row.get("number"), res)
    if profile is None:
        return finish(res)
    try:
        res.pscs = read_pscs(web, res.company_number)
    except SourceError:
        res.notes.append("persons with significant control couldn't be read, run this row again")

    last = res.last_accounts
    if last == "":
        res.status, res.note = "NOT DISCLOSED", "no accounts filed yet (new company)"
        return finish(res)
    filing = latest_accounts(web, res.company_number)
    if not filing:
        res.status = "NOT DISCLOSED" if last is None else "CHECK PDF"
        res.note = ("no accounts filed" if last is None else
                    f"accounts made up to {last} exist but weren't found in the filing history, check the company page")
        return finish(res)
    res.filed_on, res.accounts_type, res.pdf_link, xhtml = filing
    m = re.search(r"made up to (\d{1,2} \w+ \d{4})", res.accounts_type)
    if m:
        made_up = datetime.strptime(m.group(1), "%d %B %Y")
        res.period_end = made_up.date().isoformat()
        if (datetime.now() - made_up).days > STALE_DAYS:
            res.notes.append(f"latest accounts are over 2 years old ({m.group(1)})")
    reason = why_missing(res.accounts_type)

    # 1. Electronic accounts: exact figures from the iXBRL.
    if xhtml:
        try:
            r = web.get(xhtml)
        except SourceError:
            r = None
        if r is None:
            found, xhtml = None, ""             # couldn't download it: fall back to the PDF
        else:
            soup = BeautifulSoup(r.content, "html.parser")
            found = read_ixbrl(soup)
            res.cash, res.debtors = read_ixbrl_balance(soup)
            if kw:                              # the iXBRL is the whole document: search it all
                kw.scan(ixbrl_text(soup))
                res.keywords_found, res.keyword_counts = kw.result(), kw.ranked()
                kw = None                       # done, the PDF needn't be searched again
        if found:
            res.turnover, res.prior_turnover, res.currency, end, res.method, months = found
            res.period_end = end or res.period_end
            res.status = "FOUND"
            if months and not 11 <= months <= 13:
                res.notes.append(f"these accounts cover {months} months, not a normal year")
            return finish(res)
        if xhtml and reason:
            res.status, res.note = "NOT DISCLOSED", reason
            return finish(res)

    # 2. The PDF: no iXBRL, iXBRL failed to download, or full accounts with turnover untagged.
    if not res.pdf_link:
        res.status, res.note = "CHECK PDF", "the accounts filing has no downloadable document"
        return finish(res)
    r = web.get(res.pdf_link)
    if r is None:
        res.status, res.note = "CHECK PDF", "the accounts PDF couldn't be downloaded"
        return finish(res)
    pdf = read_pdf(r.content, max_pages, progress, read_scanned, kw, ocr_keywords, deadline)
    stopped = (f"{time_limit:g}-minute time limit" if pdf.timed_out else "page limit")
    if res.cash is None:
        res.cash = pdf.cash
    if res.debtors is None:
        res.debtors = pdf.debtors
    if kw and not pdf.problem:
        res.keywords_found, res.keyword_counts = kw.result(), kw.ranked()
        if pdf.scanned_skipped:
            res.notes.append("keywords not searched in scanned pages (scanned-PDF reading is turned off)")
        elif pdf.keyword_ocr_skipped:
            res.notes.append("keywords not searched in scanned pages (keyword search of scanned PDFs is turned off)")
        elif pdf.keyword_pages < pdf.pages:
            res.notes.append(f"keywords searched in {pdf.keyword_pages} of {pdf.pages} pages ({stopped})")
    if pdf.turnover is not None:
        apply_pdf(res, pdf)
    elif reason:
        res.status, res.note = "NOT DISCLOSED", reason
    elif xhtml:
        res.status, res.note = "NOT DISCLOSED", "turnover not shown in the filed accounts"
    else:
        res.status = "CHECK PDF"
        if pdf.problem:
            res.note = pdf.problem
        elif pdf.timed_out:
            res.note = (f"stopped at the {stopped} after {pdf.pages_checked} of {pdf.pages} pages, "
                        "open the PDF to check")
        elif pdf.scanned_skipped:
            res.note = "scanned PDF not read (scanned-PDF reading is turned off)"
        else:
            res.note = f"no turnover line found in {pdf.pages_checked} pages, open the PDF to check"
    if "OCR" in pdf.balance_method and res.status != "VERIFY":
        res.notes.append("cash and debtors read by OCR, check them against the PDF")
    return finish(res)


def finish(res):
    if res.currency and res.currency != "GBP":
        res.notes.append(f"figures are in {res.currency}")
    if res.notes:
        res.note = "; ".join(filter(None, [res.note] + res.notes))
    return res


def safe_check(row, web, **kw):
    """check(), but a failure only ever affects this one row."""
    try:
        return check(row, web, **kw)
    except SourceError as e:
        return Result("ERROR", note=str(e))
    except Exception as e:                      # a bug or an unexpected page: report, carry on
        traceback.print_exc(file=sys.stderr)
        return Result("ERROR", note=f"unexpected problem ({type(e).__name__}: {e}), run this row again "
                                    "or check it by hand")


# ---------------------------------------------------------------- Excel

FILL_COLOURS = {"FOUND": "C6EFCE", "VERIFY": "FFEB9C", "NOT DISCLOSED": "FCE4D6",
                "CHECK PDF": "F8CBAD", "NOT FOUND": "FFC7CE", "ERROR": "D9D9D9"}
FILLS = {k: PatternFill("solid", fgColor=v) for k, v in FILL_COLOURS.items()}
HEADER_FILL = PatternFill("solid", fgColor="7030A0")
OUT_COLS = [("Company Name", 30), ("Company Number", 17), ("Turnover Status", 16),
            ('Turnover or "Revenue"', 22), ('Prior Year "Turnover" or "Revenue"', 33), ("Cash", 16),
            ("Debtors", 16), ("Persons with significant control", 40), ("linkedin search", 45),
            ("Keyword Count", 15), ("Keywords Found", 100), ("Notes", 60), ("Latest Accounts", 46),
            ("Accounts PDF", 13), ("Companies House Page", 22)]
MONEY_COLS = ('Turnover or "Revenue"', 'Prior Year "Turnover" or "Revenue"', "Cash", "Debtors")
MONEY = '_-[$£-809]* #,##0.00_-;\\-[$£-809]* #,##0.00_-;_-[$£-809]* "-"??_-;_-@_-'   # £ accounting
PLAIN_MONEY = '#,##0;(#,##0)'                       # for accounts in another currency
MEANINGS = {
    "FOUND": "Read exactly from the filed accounts.",
    "VERIFY": "Read from a scanned PDF with OCR. Check the figure against the PDF (page given).",
    "NOT DISCLOSED": "The company didn't file a turnover figure. The Notes column says why.",
    "CHECK PDF": "Couldn't read a figure automatically. Open the linked PDF.",
    "NOT FOUND": "Couldn't match the company on the register. Add or check the company number.",
    "ERROR": "Companies House couldn't be reached for this row. Run it again later.",
}
LEGAL_ENDING = re.compile(r"\s+(limited|ltd\.?|plc|p\.l\.c\.?|llp|public limited company)$", re.I)
TITLE = re.compile(r"^(mr|mrs|ms|miss|mx|dr|sir|dame|lord|lady|prof|professor|rev)\.?\s+", re.I)


def keyword_count(res):
    """How many different keywords were found; None when nothing was searched."""
    return None if res.keywords_found is None else len(res.keywords_found)


def keyword_cell(res):
    """'foreign exchange (14), hedging (3)', most mentioned first; None when nothing was searched."""
    if res.keywords_found is None:
        return None
    if not res.keywords_found:
        return "none found"
    counts = res.keyword_counts or {w: 1 for w in res.keywords_found}
    return ", ".join(f"{w} ({n})" for w, n in counts.items())


def psc_cell(res):
    if res.pscs is None:
        return None
    return "; ".join(n for n, _corp in res.pscs) if res.pscs else "none registered"


def linkedin_search(res):
    """(cell text, Google search URL) for the company plus its first individual PSC (else its
    first PSC) plus 'LinkedIn'; (None, None) when the company wasn't found."""
    company = LEGAL_ENDING.sub("", (res.company_name or "").strip())
    if not company:
        return None, None
    people = res.pscs or []
    person = next((n for n, corporate in people if not corporate), people[0][0] if people else "")
    query = " ".join(filter(None, [company, TITLE.sub("", person), "LinkedIn"]))
    return " + ".join(filter(None, [company, person, "LinkedIn"])), "https://www.google.com/search?q=" + quote_plus(query)


def find_col(ws, header, header_row=1):
    if not header:
        return None
    for cell in ws[header_row]:
        if cell.value is not None and str(cell.value).strip().lower() == header.strip().lower():
            return cell.column
    found = [str(c.value) for c in ws[header_row] if c.value is not None]
    raise KeyError(f"Column '{header}' not found in row {header_row}. Headers there are: {found}")


def new_output():
    """A workbook with a 'Results' sheet and its header row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    for i, (h, w) in enumerate(OUT_COLS, 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return wb, ws


def write_row(ws, r, row, res):
    """One output row. row is the input {'name', 'number'}, used when the company wasn't found."""
    search_text, search_url = linkedin_search(res)
    values = {"Company Name": res.company_name or str(row.get("name") or "").strip() or None,
              "Company Number": res.company_number or clean_number(row.get("number")) or None,
              "Turnover Status": res.status, 'Turnover or "Revenue"': res.turnover,
              'Prior Year "Turnover" or "Revenue"': res.prior_turnover, "Cash": res.cash, "Debtors": res.debtors,
              "Persons with significant control": psc_cell(res), "linkedin search": search_text,
              "Keyword Count": keyword_count(res), "Keywords Found": keyword_cell(res), "Notes": res.note,
              "Latest Accounts": res.accounts_type, "Accounts PDF": "Open PDF" if res.pdf_link else None,
              "Companies House Page": "Open page" if res.company_number else None}
    links = {"linkedin search": search_url, "Accounts PDF": res.pdf_link, "Companies House Page": res.company_link}
    for i, (h, _w) in enumerate(OUT_COLS, 1):
        v = values[h]
        cell = ws.cell(row=r, column=i, value=v if v != "" else None)
        cell.alignment = Alignment(vertical="top", wrap_text=h in ("Keywords Found", "Notes"))
        if h in MONEY_COLS and v is not None:
            cell.number_format = MONEY if res.currency in ("", "GBP") else PLAIN_MONEY
        if links.get(h):
            cell.hyperlink, cell.font = links[h], Font(color="0563C1", underline="single")
    ws.cell(row=r, column=3).fill = FILLS.get(res.status, FILLS["ERROR"])


def process_sheet(src, cols, out_ws, web=None, max_pages=None, on_row=None, on_detail=None,
                  checkpoint=None, header_row=1, read_scanned=True, keywords=None, ocr_keywords=True,
                  time_limit=0):
    """Look up every company in src (rows below header_row) and write one row per company to
    out_ws (made by new_output). cols maps 'name' and 'number' to src column indexes."""
    web = web or Web()
    cache, counts = {}, {}
    rows = [r for r in range(header_row + 1, src.max_row + 1)
            if any(src.cell(row=r, column=c).value not in (None, "") for c in cols.values() if c)]
    for done, r in enumerate(rows, 1):
        row = {k: src.cell(row=r, column=c).value for k, c in cols.items() if c}
        key = (str(row.get("name") or "").strip().lower(), clean_number(row.get("number")))
        if key not in cache:
            label = str(row.get("name") or row.get("number"))
            detail = (lambda msg, label=label: on_detail(f"{label}: {msg}")) if on_detail else None
            cache[key] = safe_check(row, web, max_pages=max_pages, progress=detail, read_scanned=read_scanned,
                                    keywords=keywords, ocr_keywords=ocr_keywords, time_limit=time_limit)
        res = cache[key]
        write_row(out_ws, done + 1, row, res)
        counts[res.status] = counts.get(res.status, 0) + 1
        if res.keywords_found:
            counts["_with_keywords"] = counts.get("_with_keywords", 0) + 1
        if on_row:
            on_row(done, len(rows), str(row.get("name") or row.get("number")), res)
        if checkpoint and done % 10 == 0:
            checkpoint()
    out_ws.auto_filter.ref = f"A1:{get_column_letter(len(OUT_COLS))}{max(len(rows), 1) + 1}"
    return counts


def add_summary(wb, counts, source_name="", keywords=None):
    """A 'Summary' sheet: counts per status, what each status means, when it ran, and which
    keywords were searched for."""
    title = "Summary"
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 80
    ws["A1"] = "Automation Tool results"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Run on {datetime.now():%d %B %Y at %H:%M}" + (f" for {source_name}" if source_name else "")
    ws["A3"] = "Source: Companies House (find-and-update.company-information.service.gov.uk)"
    for i, h in enumerate(["Status", "Companies", "Meaning"]):
        c = ws.cell(row=5, column=i + 1, value=h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
    for r, status in enumerate(STATUSES, 6):
        ws.cell(row=r, column=1, value=status).fill = FILLS[status]
        ws.cell(row=r, column=2, value=counts.get(status, 0))
        ws.cell(row=r, column=3, value=MEANINGS[status])
    total = len(STATUSES) + 6
    ws.cell(row=total, column=1, value="Total").font = Font(bold=True)
    ws.cell(row=total, column=2, value=sum(v for k, v in counts.items() if k in STATUSES)).font = Font(bold=True)
    if keywords:
        r = total + 2
        ws.cell(row=r, column=1, value="Keywords").font = Font(bold=True)
        ws.cell(row=r, column=2, value=counts.get("_with_keywords", 0))
        ws.cell(row=r, column=3, value="companies whose accounts mention at least one keyword")
        ws.cell(row=r + 1, column=3, value=f"Searched for {len(keywords)}: " + ", ".join(keywords))
        ws.cell(row=r + 1, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input")
    ap.add_argument("--name-col", help="column with company names")
    ap.add_argument("--number-col", help="column with company numbers")
    ap.add_argument("--sheet")
    ap.add_argument("--header-row", type=int, default=1, help="row the column headers are in")
    ap.add_argument("--max-pages", type=int, default=0, help="pages read per PDF (0 = no limit)")
    ap.add_argument("--time-limit", type=float, default=0,
                    help="minutes allowed per company for reading its PDF (0 = no limit)")
    ap.add_argument("--no-ocr", action="store_true", help="don't read scanned PDFs (much faster)")
    ap.add_argument("--keywords", help="words or phrases to look for in the accounts, comma-separated")
    ap.add_argument("--keywords-file", help="text file of keywords, one per line")
    ap.add_argument("--no-keyword-ocr", action="store_true",
                    help="don't search scanned PDFs for keywords (they're read page by page, which is slow)")
    ap.add_argument("--output")
    args = ap.parse_args()
    if not (args.name_col or args.number_col):
        sys.exit("Give --name-col, --number-col or both.")
    keywords = parse_keywords(args.keywords) if args.keywords else []
    if args.keywords_file:
        with open(args.keywords_file, encoding="utf-8") as f:
            keywords += parse_keywords(f.read())

    out_path = args.output or re.sub(r"\.(xlsx|xlsm|xls|csv)$", "", args.input, flags=re.I) + "_results.xlsx"
    src_wb = load_workbook(args.input, data_only=True)
    src = src_wb[args.sheet] if args.sheet else src_wb.worksheets[0]
    try:
        cols = {"name": find_col(src, args.name_col, args.header_row),
                "number": find_col(src, args.number_col, args.header_row)}
    except KeyError as e:
        sys.exit(e.args[0])

    def show(i, n, name, res):
        amount = f"{res.turnover:,.0f}" if res.turnover is not None else ""
        kws = f"  [{len(res.keywords_found)} keywords]" if res.keywords_found else ""
        print(f"[{i}/{n}] {name[:40]:40} {res.status:13} {amount:>16}{kws}  {res.note[:60]}", flush=True)

    out_wb, out_ws = new_output()
    counts = process_sheet(src, cols, out_ws, time_limit=args.time_limit, max_pages=args.max_pages or None,
                           on_row=show, header_row=args.header_row, read_scanned=not args.no_ocr,
                           keywords=keywords or None, ocr_keywords=not args.no_keyword_ocr,
                           checkpoint=lambda: out_wb.save(out_path))
    add_summary(out_wb, counts, args.input, keywords)
    out_wb.save(out_path)
    print("\n" + ", ".join(f"{k}: {v}" for k, v in sorted(counts.items()) if k in STATUSES))
    print(f"Saved to {out_path}")


if __name__ == "__main__":
    main()
